"""Level-2 exit gate, end to end over HTTP on the real app (no stubs).

"Invoice file goes from upload to validated export without manual steps."
Output: 22 invoice rows, 313 line items, gross total 154,292 matched.
Rollback restores the original exactly.
"""

from __future__ import annotations

import asyncio
import hashlib
import io
import json
import runpy
import zipfile
from pathlib import Path

import openpyxl
import pytest
from alembic import command
from alembic.config import Config
from fastapi.testclient import TestClient

BACKEND = Path(__file__).resolve().parents[2]
REFERENCE = BACKEND.parent / "data" / "reference" / "VendorInvoices_uncleaned.xlsx"


@pytest.fixture(scope="module")
def client() -> TestClient:
    cfg = Config(str(BACKEND / "alembic.ini"))
    cfg.set_main_option("script_location", str(BACKEND / "migrations"))
    command.downgrade(cfg, "base")  # the session DB may hold create_all() tables from other tests
    command.upgrade(cfg, "head")
    seed = runpy.run_path(str(BACKEND / "scripts" / "seed_demo.py"))
    asyncio.run(seed["main"]())

    from planner.main import app

    c = TestClient(app)
    r = c.post("/api/v1/auth/login", json={"email": "admin@example.com", "password": "Admin123456!"})
    assert r.status_code == 200, r.text
    c.headers["Authorization"] = f"Bearer {r.json()['accessToken']}"
    return c


def _ok(r, *codes: int):
    assert r.status_code in (codes or (200, 201, 202)), f"{r.request.method} {r.request.url}: {r.status_code} {r.text}"
    return r.json()


def test_exit_gate__reference_invoice_file__upload_to_validated_export_and_rollback(
    client: TestClient,
) -> None:
    api = "/api/v1"
    original_bytes = REFERENCE.read_bytes()

    # ingest + profile + infer run on upload
    ds = _ok(client.post(f"{api}/datasets", files={"file": (REFERENCE.name, original_bytes)}))
    ds = _ok(client.get(f"{api}/datasets/{ds['id']}"))
    assert (ds["status"], ds["rowCount"], ds["columnCount"]) == ("profiled", 22, 14)

    rules = _ok(client.get(f"{api}/datasets/{ds['id']}/rules"))["items"]
    supplier = next(r for r in rules if r["ruleType"] == "entity_group" and r["columns"] == ["supplier_name"])
    groups = supplier["expression"]["groups"]
    assert (sum(len(g["variants"]) for g in groups), len(groups)) == (7, 5)  # supplier-variants-7-to-5

    # plan with per-step loss estimates, every step decided, then approve (runs tests + execute)
    plan = _ok(client.post(f"{api}/datasets/{ds['id']}/plans"))
    ops = [s["operation"] for s in plan["steps"]]
    assert {"drop_column", "replace_value", "expand_nested"} <= set(ops)
    assert all(s["estimatedLoss"] is not None for s in plan["steps"])
    assert plan["totalEstimatedLoss"] < plan["lossThreshold"]
    for s in plan["steps"]:
        _ok(client.patch(f"{api}/plans/{plan['id']}/steps/{s['id']}", json={"decision": "accepted"}))
    _ok(client.post(f"{api}/plans/{plan['id']}/approve"))

    # verify: generated tests pass and the numbers reconcile against the source
    v = _ok(client.get(f"{api}/plans/{plan['id']}/validation"))
    assert v["passed"] is True
    assert all(t["latestRun"]["result"] == "passed" for t in v["testCases"])
    rec = {r["checkName"]: r for r in v["reconciliation"]}
    assert all(r["ok"] for r in rec.values()), rec
    assert rec["row_count"]["outputValue"] == "22"
    assert rec["row_count:LineItems"]["outputValue"] == "313"
    gross = next(r for name, r in rec.items() if name.startswith("gross_total:"))
    assert round(float(gross["sourceValue"])) == round(float(gross["outputValue"])) == 154292

    # validated export: every table, every format
    def download(fmt: str) -> bytes:
        url = _ok(client.post(f"{api}/plans/{plan['id']}/exports", json={"format": fmt}))["downloadUrl"]
        r = client.get(url)
        assert r.status_code == 200
        return r.content

    wb = openpyxl.load_workbook(io.BytesIO(download("xlsx")))
    assert {ws.title: ws.max_row - 1 for ws in wb} == {"main": 22, "LineItems": 313}
    sheet = wb["main"]
    header = [c.value for c in sheet[1]]
    suppliers = {row[header.index("supplier_name")] for row in sheet.iter_rows(min_row=2, values_only=True)}
    assert len(suppliers) == 5
    assert "total_price" not in header and "customer_vat_number" not in header
    gross_col = header.index("subtotal_with_vat")
    assert round(sum(r[gross_col] for r in sheet.iter_rows(min_row=2, values_only=True))) == 154292

    z = zipfile.ZipFile(io.BytesIO(download("csv")))
    assert sorted(z.namelist()) == ["LineItems.csv", "main.csv"]
    pipeline = json.loads(download("pipeline"))
    assert [s["operation"] for s in pipeline["steps"]] == ops

    # rollback to the original: byte-identical snapshot, raw upload untouched
    versions = _ok(client.get(f"{api}/plans/{plan['id']}/versions"))["items"]
    last = max(x["versionNo"] for x in versions)
    _ok(client.post(f"{api}/plans/{plan['id']}/rollback", json={"toVersion": 0, "reason": "exit gate: restore the original"}))
    versions = _ok(client.get(f"{api}/plans/{plan['id']}/versions"))["items"]
    assert [x["versionNo"] for x in versions if x["isCurrent"]] == [last + 1]

    from planner.core.config import settings

    root = Path(settings.storage_local_root)
    snap = root / "snapshots" / plan["id"]
    assert hashlib.sha256((snap / f"v{last + 1}.parquet").read_bytes()).digest() == hashlib.sha256(
        (snap / "v0.parquet").read_bytes()
    ).digest()
    assert (root / ds["rawObjectKey"]).read_bytes() == original_bytes  # FR-004 original immutable

    # rolled-back state is not validated, so export is withheld (FR-043)
    r = client.post(f"{api}/plans/{plan['id']}/exports", json={"format": "xlsx"})
    assert r.status_code == 409

    # every stage of the loop is in the audit trail (FR-051)
    events = {e["eventType"] for e in _ok(client.get(f"{api}/audit-events", params={"pageSize": 100}))["items"]}
    assert {
        "dataset.upload", "dataset.profiled", "rules.inferred", "plan.created", "plan.approved",
        "plan.executed", "validation.completed", "plan.exported", "rollback.completed",
    } <= events
